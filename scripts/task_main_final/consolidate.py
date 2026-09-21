from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import shutil
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence


ROOT = Path(__file__).resolve().parents[2]
FORMAL = ROOT / "results/task_main/formal"
DIAGNOSTICS = ROOT / "results/task_main/formal_diagnostics_v1"
D1 = ROOT / "results/task_main/stage_d1_sensitivity/run_20260917_01"
D2 = ROOT / "results/task_main/stage_d2_move_ablation/run_20260918_02"
DEFAULT_OUTPUT = ROOT / "results/task_main/final_results"
AUDIT_REPORT = ROOT / "docs/analysis/task_main_final_results_consolidation.md"

PROTOCOL = "TASK_MAIN_V1_STAGE_C"
FORMAL_RUN_ID = "formal:5ddeca8e0c768e04"
FORMAL_MANIFEST_SHA256 = "7499b1e08fe0f4b7a8bbd72c83b120c16ebfc6a647d797095ca469c056a4debf"
FORMAL_SOURCE_PROVENANCE_SHA256 = "25ddfae1be89148d8f30a39b03fd0b4764b0703041d4a7be4e00d5b69f4eb318"
TRACE_SHA256 = {
    "conversation": "b8cbb061a85206d729d91cdc2981f43c9e0d99209dce588d3af5f7934408b9df",
    "toolagent": "48a2db1a13d3bc05e6330140c64f604ba366df20d3c9e128b5c35a01c1fa5f71",
}
MAIN_EVAL = (1_500_000, 2_700_000)
W_SENS_EVAL = (600_000, 1_500_000)
CAPACITY_PER_POD = 585

LINE_STRATEGY = {
    1: "R_AFF",
    2: "R_REQ_KV",
    3: "Future-Demand Oracle",
    4: "Persistence h=60s K=10",
    5: "Persistence h=300s K=10",
    6: "Recency q=0.9",
    7: "Persistence + Cost-aware",
}

FINAL_LONG_COLUMNS = [
    "canonical_case_id", "workload", "family", "referenced_by_families",
    "source_kind", "source_run_id", "source_artifact", "protocol_version",
    "trace_sha256", "validation_status", "deterministic", "line_id", "case_id",
    "routing", "policy", "action", "N", "capacity_pages_per_pod",
    "total_cluster_capacity_pages", "theta", "W_seconds", "h_seconds", "K", "q",
    "evaluation_start_ms", "evaluation_end_ms", "evaluation_label",
    "saved_tokens", "saved_tokens_million", "weighted_saved_tokens",
    "weighted_saved_tokens_million", "request_hit_rate", "token_hit_rate", "skew_ratio",
    "eval_wire_bytes", "eval_wire_tb_decimal", "full_wire_bytes", "full_wire_tb_decimal",
    "eval_reactive_wire_bytes", "eval_proactive_wire_bytes",
    "full_reactive_wire_bytes", "full_proactive_wire_bytes",
    "transfer_count_evaluation", "transfer_count_full_run", "unused_transfer_wire_ratio",
    "waste_ratio_original_name", "proactive_actions_evaluation",
    "proactive_actions_full_run", "reactive_actions_evaluation", "reactive_actions_full_run",
    "ordinary_lru_evictions_evaluation", "ordinary_lru_evictions_full_run",
    "ordinary_cache_turnover_pages_evaluation", "ordinary_cache_turnover_pages_full_run",
    "cache_eviction_count_all_causes_evaluation", "cache_eviction_count_all_causes_full_run",
    "cache_admission_skips_evaluation", "cache_admission_skips_full_run",
    "no_capacity_candidate_skips_evaluation", "no_capacity_candidate_skips_full_run",
    "move_actions", "source_released_pages", "source_released_bytes",
    "release_fraction_mean", "release_fraction_median", "zero_release_actions",
    "partial_release_actions", "full_release_actions", "chain_depth_mean",
    "chain_depth_median", "chain_depth_p90", "new_page_observed_use_ratio",
    "positive_candidate_count_evaluation", "shortlist_count_evaluation",
    "capacity_metric_scope", "notes",
]


def load_json(path: Path) -> Any:
    with path.open("r", encoding="utf-8") as handle:
        return json.load(handle)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def require(condition: bool, message: str) -> None:
    if not condition:
        raise RuntimeError(message)


def all_true(values: Mapping[str, Any]) -> bool:
    return all(value is True for value in values.values())


def equal_number(left: Any, right: Any, *, tolerance: float = 1e-9) -> bool:
    if left is None or right is None:
        return left is right
    if isinstance(left, (int, float)) and isinstance(right, (int, float)):
        return math.isclose(float(left), float(right), rel_tol=tolerance, abs_tol=tolerance)
    return left == right


def percent_delta(value: float | int | None, reference: float | int | None) -> float | None:
    if value is None or reference in (None, 0):
        return None
    return (float(value) - float(reference)) / float(reference) * 100.0


def cohort_label(start: int, end: int) -> str:
    if (start, end) == MAIN_EVAL:
        return "MAIN_25_45"
    if (start, end) == W_SENS_EVAL:
        return "W_SENS_10_25"
    return f"EVAL_{start}_{end}"


def slug_number(value: Any) -> str:
    if value is None:
        return "na"
    number = float(value)
    if number.is_integer():
        return str(int(number))
    return str(number).replace(".", "p")


def slug_text(value: Any) -> str:
    return str(value).strip().lower().replace("+", "plus").replace("_", "-").replace(" ", "-")


def behavior_key(row: Mapping[str, Any]) -> tuple[Any, ...]:
    return (
        row["workload"], row["routing"], row["policy"], row["action"], row["N"],
        float(row["theta"]), float(row["W_seconds"]), float(row["h_seconds"]),
        int(row["K"]), float(row["q"]), row["evaluation_start_ms"], row["evaluation_end_ms"],
    )


def canonical_id(row: Mapping[str, Any]) -> str:
    return "__".join([
        slug_text(row["workload"]), slug_text(row["family"]), slug_text(row["routing"]),
        slug_text(row["policy"]), slug_text(row["action"]), f"n{row['N']}",
        f"theta{slug_number(row['theta'])}", f"w{slug_number(row['W_seconds'])}",
        f"h{slug_number(row['h_seconds'])}", f"k{row['K']}", f"q{slug_number(row['q'])}",
        f"eval{row['evaluation_start_ms']}-{row['evaluation_end_ms']}",
    ])


def normalize_row(row: dict[str, Any]) -> dict[str, Any]:
    row["saved_tokens_million"] = row["saved_tokens"] / 1e6
    row["weighted_saved_tokens_million"] = row["weighted_saved_tokens"] / 1e6
    row["eval_wire_tb_decimal"] = row["eval_wire_bytes"] / 1e12
    row["full_wire_tb_decimal"] = row["full_wire_bytes"] / 1e12
    row["total_cluster_capacity_pages"] = row["N"] * row["capacity_pages_per_pod"]
    row["evaluation_label"] = cohort_label(row["evaluation_start_ms"], row["evaluation_end_ms"])
    row["canonical_case_id"] = canonical_id(row)
    for key in FINAL_LONG_COLUMNS:
        row.setdefault(key, None)
    return row


def write_json(path: Path, rows: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        json.dump(rows, handle, ensure_ascii=False, indent=2, sort_keys=True, allow_nan=False)
        handle.write("\n")


def csv_value(value: Any) -> Any:
    if value is None:
        return "NA"
    if isinstance(value, bool):
        return "true" if value else "false"
    return value


def write_csv(path: Path, rows: Sequence[Mapping[str, Any]], columns: Sequence[str] | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if columns is None:
        columns = list(rows[0]) if rows else []
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(columns), extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({column: csv_value(row.get(column)) for column in columns})


def write_pair(output: Path, stem: str, rows: Sequence[Mapping[str, Any]], columns: Sequence[str] | None = None) -> None:
    write_json(output / f"{stem}.json", list(rows))
    write_csv(output / f"{stem}.csv", rows, columns)


def network_lookup(rows: Sequence[Mapping[str, Any]]) -> dict[tuple[Any, ...], Mapping[str, Any]]:
    return {
        (r["workload"], r["line_id"], r["phase"], r["type"]): r
        for r in rows
    }


def source_gate() -> dict[str, Any]:
    formal_manifest_path = FORMAL / "validation_manifest.json"
    formal_manifest = load_json(formal_manifest_path)
    formal_checkpoint = load_json(FORMAL / "checkpoint.json")
    formal_progress = load_json(FORMAL / "progress.json")
    require(formal_manifest["status"] == "PASS", "Formal manifest is not PASS")
    require(len(formal_manifest["cases"]) == 14, "Formal manifest does not contain 14 cases")
    require(formal_manifest["completed_replays"] == 28, "Formal does not contain 28 replays")
    require(all(c["status"] == "PASS" and c["deterministic"] for c in formal_manifest["cases"]),
            "Formal cases are not all PASS and deterministic")
    require(formal_checkpoint["status"] == "PASS" and len(formal_checkpoint["completed"]) == 28,
            "Formal checkpoint is not 28/28 PASS")
    require(formal_progress["status"] == "PASS" and formal_progress["completed_replays"] == 28,
            "Formal progress is not PASS")
    require(sha256(formal_manifest_path) == FORMAL_MANIFEST_SHA256, "Formal manifest SHA-256 changed")
    require(sha256(FORMAL / "source_provenance.json") == FORMAL_SOURCE_PROVENANCE_SHA256,
            "Formal source provenance SHA-256 changed")

    diagnostics_manifest = load_json(DIAGNOSTICS / "diagnostics_manifest.json")
    diagnostics_validation = load_json(DIAGNOSTICS / "diagnostics_validation.json")
    run_identity = load_json(DIAGNOSTICS / "run_identity.json")
    require(diagnostics_manifest["status"] == "PASS", "Diagnostics manifest is not PASS")
    require(diagnostics_validation["status"] == "PASS" and all_true(diagnostics_validation["checks"]),
            "Diagnostics validation is not PASS")
    require(run_identity["formal_run_id"] == FORMAL_RUN_ID and all_true(run_identity["checks"]),
            "Diagnostics run identity is invalid")
    require(diagnostics_manifest["formal_artifacts_read_only"] is True
            and diagnostics_manifest["formal_14_rerun"] is False
            and diagnostics_manifest["parameter_changes"] is False,
            "Diagnostics violated read-only constraints")
    for name, expected in diagnostics_manifest["artifact_sha256"].items():
        require(sha256(DIAGNOSTICS / name) == expected, f"Diagnostics artifact changed: {name}")

    d1_manifest = load_json(D1 / "validation_manifest.json")
    d1_checkpoint = load_json(D1 / "checkpoint.json")
    d1_progress = load_json(D1 / "progress.json")
    require(d1_manifest["status"] == "PASS" and d1_manifest["case_count"] == 32
            and d1_manifest["replay_count"] == 64, "D1 manifest is not 32-case PASS")
    require(all(c["status"] == "PASS" and c["deterministic"] for c in d1_manifest["cases"]),
            "D1 cases are not deterministic PASS")
    require(d1_checkpoint["status"] == "PASS" and len(d1_checkpoint["completed"]) == 64,
            "D1 checkpoint is not 64/64 PASS")
    require(d1_progress["status"] == "PASS" and d1_progress["completed_replays"] == 64,
            "D1 progress is not PASS")
    d1_ref = d1_manifest["formal_reference"]
    require(d1_ref["status"] == "PASS" and all_true(d1_ref["checks"]), "D1 formal reference gate failed")
    require(d1_ref["validation_manifest_sha256"] == FORMAL_MANIFEST_SHA256
            and d1_ref["source_provenance_sha256"] == FORMAL_SOURCE_PROVENANCE_SHA256,
            "D1 formal provenance differs")

    d2_manifest = load_json(D2 / "validation_manifest.json")
    d2_checkpoint = load_json(D2 / "checkpoint.json")
    d2_progress = load_json(D2 / "progress.json")
    require(d2_manifest["status"] == "PASS" and d2_manifest["case_count"] == 4
            and d2_manifest["replay_count"] == 8, "D2 manifest is not 4-case PASS")
    require(all(c["deterministic"] for c in d2_manifest["cases"]), "D2 cases are not deterministic")
    require(d2_checkpoint["status"] == "PASS" and len(d2_checkpoint["completed"]) == 8,
            "D2 checkpoint is not 8/8 PASS")
    require(d2_progress["status"] == "PASS" and d2_progress["completed_replays"] == 8,
            "D2 progress is not PASS")
    d2_ref = d2_manifest["copy_reference"]
    require(d2_ref["status"] == "PASS" and all_true(d2_ref["checks"]), "D2 COPY reference gate failed")
    require(d2_ref["formal_manifest_sha256"] == FORMAL_MANIFEST_SHA256
            and d2_ref["formal_source_provenance_sha256"] == FORMAL_SOURCE_PROVENANCE_SHA256,
            "D2 formal provenance differs")

    for manifest in (d1_manifest, d2_manifest):
        for case in manifest["cases"]:
            config = case["config"]
            require(config["protocol_version"] == PROTOCOL, "Protocol version mismatch")
            require(config["trace_sha256"] == TRACE_SHA256[config["workload"]], "Trace SHA mismatch")

    return {
        "formal_manifest": formal_manifest,
        "diagnostics_manifest": diagnostics_manifest,
        "diagnostics_validation": diagnostics_validation,
        "run_identity": run_identity,
        "d1_manifest": d1_manifest,
        "d2_manifest": d2_manifest,
        "source_sha256": {
            "formal_validation_manifest": sha256(formal_manifest_path),
            "formal_source_provenance": sha256(FORMAL / "source_provenance.json"),
            "diagnostics_manifest": sha256(DIAGNOSTICS / "diagnostics_manifest.json"),
            "d1_validation_manifest": sha256(D1 / "validation_manifest.json"),
            "d2_validation_manifest": sha256(D2 / "validation_manifest.json"),
        },
    }


def formal_rows(gate: Mapping[str, Any], d2_table: Sequence[Mapping[str, Any]]) -> list[dict[str, Any]]:
    strategy = load_json(FORMAL / "tables/strategy_workload.json")
    transfer = network_lookup(load_json(FORMAL / "tables/transfer_cost.json"))
    churn_rows = load_json(DIAGNOSTICS / "cache_churn_summary.json")
    churn = {(r["workload"], r["line_id"]): r for r in churn_rows}
    d2_copy = {(r["workload"], r["trigger"]): r for r in d2_table if r["action"] == "COPY"}
    output: list[dict[str, Any]] = []
    for source in strategy:
        workload, line = source["workload"], source["line_id"]
        summary = load_json(FORMAL / f"{workload}_line_{line}/run_1/summary.json")
        config = summary["provenance"]["config"]
        require(config["trace_sha256"] == TRACE_SHA256[workload], "Formal trace identity mismatch")
        require((config["evaluation_start_ms"], config["evaluation_end_ms"]) == MAIN_EVAL,
                "Formal cohort mismatch")
        net = lambda phase, kind: transfer[(workload, line, phase, kind)]
        row: dict[str, Any] = {
            "workload": workload, "family": "MAIN", "referenced_by_families": "MAIN",
            "source_kind": "FORMAL_MAIN", "source_run_id": FORMAL_RUN_ID,
            "source_artifact": "results/task_main/formal/tables/strategy_workload.json",
            "protocol_version": source["protocol_version"], "trace_sha256": config["trace_sha256"],
            "validation_status": "PASS", "deterministic": True, "line_id": line,
            "case_id": f"line_{line}", "routing": source["routing"], "policy": source["policy"],
            "action": source["action"], "N": source["N"],
            "capacity_pages_per_pod": config["capacity_pages"], "theta": source["theta"],
            "W_seconds": source["W"], "h_seconds": source["h"], "K": source["K"], "q": source["q"],
            "evaluation_start_ms": config["evaluation_start_ms"],
            "evaluation_end_ms": config["evaluation_end_ms"],
            "saved_tokens": source["saved_prefill_tokens"],
            "weighted_saved_tokens": source["weighted_saved_tokens"],
            "request_hit_rate": source["request_hit_rate"], "token_hit_rate": source["token_hit_rate"],
            "skew_ratio": source["skew_ratio"],
            "eval_wire_bytes": net("EVAL", "TOTAL")["wire_bytes"],
            "full_wire_bytes": net("FULL_RUN", "TOTAL")["wire_bytes"],
            "eval_reactive_wire_bytes": net("EVAL", "REACTIVE")["wire_bytes"],
            "eval_proactive_wire_bytes": net("EVAL", "PROACTIVE")["wire_bytes"],
            "full_reactive_wire_bytes": net("FULL_RUN", "REACTIVE")["wire_bytes"],
            "full_proactive_wire_bytes": net("FULL_RUN", "PROACTIVE")["wire_bytes"],
            "transfer_count_evaluation": net("EVAL", "TOTAL")["transfer_started_count"],
            "transfer_count_full_run": net("FULL_RUN", "TOTAL")["transfer_started_count"],
            "unused_transfer_wire_ratio": source["wasted_copy_ratio"],
            "waste_ratio_original_name": "wasted_copy_ratio" if source["wasted_copy_ratio"] is not None else None,
            "proactive_actions_evaluation": net("EVAL", "PROACTIVE")["transfer_started_count"],
            "proactive_actions_full_run": net("FULL_RUN", "PROACTIVE")["transfer_started_count"],
            "reactive_actions_evaluation": net("EVAL", "REACTIVE")["transfer_started_count"],
            "reactive_actions_full_run": net("FULL_RUN", "REACTIVE")["transfer_started_count"],
            "capacity_metric_scope": "NA_NOT_EMITTED_BY_CANONICAL_SUMMARY",
            "notes": "Canonical Formal Main result",
        }
        diagnostic = churn.get((workload, line))
        if diagnostic is None and line == 7:
            diagnostic = churn.get((workload, 5))
            row["notes"] += "; capacity diagnostics inherited via validated Line5/Line7 execution equivalence"
        if diagnostic is not None:
            row.update({
                "cache_eviction_count_all_causes_evaluation": diagnostic["evaluation_cache_eviction_records_all_causes"],
                "cache_eviction_count_all_causes_full_run": diagnostic["total_cache_eviction_records_all_causes"],
                "cache_admission_skips_evaluation": diagnostic["cache_admission_skip_evaluation"],
                "cache_admission_skips_full_run": diagnostic["cache_admission_skip_full_run"],
                "no_capacity_candidate_skips_evaluation": diagnostic["no_capacity_candidate_skips_evaluation"],
                "no_capacity_candidate_skips_full_run": diagnostic["no_capacity_candidate_skips_full_run"],
                "capacity_metric_scope": "FORMAL_DIAGNOSTICS_ALL_CAUSES",
            })
        if line in (3, 5, 7):
            trigger = "FUTURE_DEMAND" if line == 3 else "PERSISTENCE"
            ref = d2_copy[(workload, trigger)]
            row.update({
                "ordinary_lru_evictions_evaluation": ref["ordinary_lru_evictions_evaluation"],
                "ordinary_lru_evictions_full_run": ref["ordinary_lru_evictions_full_run"],
                "ordinary_cache_turnover_pages_evaluation": ref["ordinary_cache_turnover_pages_evaluation"],
                "ordinary_cache_turnover_pages_full_run": ref["ordinary_cache_turnover_pages_full_run"],
                "cache_admission_skips_evaluation": ref["cache_admission_skips_evaluation"],
                "cache_admission_skips_full_run": ref["cache_admission_skips_full_run"],
                "no_capacity_candidate_skips_evaluation": ref["no_capacity_candidate_skips_evaluation"],
                "no_capacity_candidate_skips_full_run": ref["no_capacity_candidate_skips_full_run"],
                "capacity_metric_scope": "D2_READ_ONLY_FORMAL_REFERENCE_ORDINARY_LRU",
            })
            if line == 7:
                row["notes"] += "; ordinary-LRU fields inherited via validated Line5/Line7 execution equivalence"
        output.append(normalize_row(row))
    return output


def d1_rows(gate: Mapping[str, Any]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    source_rows = load_json(D1 / "tables/all_sensitivity_long.json")
    manifest_cases = {(c["workload"], c["case_id"]): c for c in gate["d1_manifest"]["cases"]}
    output: list[dict[str, Any]] = []
    for source in source_rows:
        if source["source_kind"] != "STAGE_D1_REPLAY":
            continue
        case = manifest_cases[(source["workload"], source["case_id"])]
        config = case["config"]
        require(case["deterministic"] and config["protocol_version"] == source["protocol_version"],
                "D1 case identity mismatch")
        row = {
            "workload": source["workload"], "family": source["family"],
            "referenced_by_families": source["family"], "source_kind": "STAGE_D1_REPLAY",
            "source_run_id": source["source_run_id"],
            "source_artifact": "results/task_main/stage_d1_sensitivity/run_20260917_01/tables/all_sensitivity_long.json",
            "protocol_version": source["protocol_version"], "trace_sha256": config["trace_sha256"],
            "validation_status": "PASS", "deterministic": case["deterministic"], "line_id": None,
            "case_id": source["case_id"], "routing": source["routing"], "policy": source["policy"],
            "action": source["action"], "N": source["N"],
            "capacity_pages_per_pod": source["capacity_pages_per_pod"], "theta": source["theta"],
            "W_seconds": source["W"], "h_seconds": source["h"], "K": source["K"], "q": source["q"],
            "evaluation_start_ms": source["evaluation_start"], "evaluation_end_ms": source["evaluation_end"],
            "saved_tokens": source["saved_tokens"], "weighted_saved_tokens": source["weighted_saved_tokens"],
            "request_hit_rate": source["request_hit_rate"], "token_hit_rate": source["token_hit_rate"],
            "skew_ratio": source["skew_ratio"], "eval_wire_bytes": source["eval_wire_bytes"],
            "full_wire_bytes": source["full_wire_bytes"],
            "eval_reactive_wire_bytes": source["eval_reactive_wire_bytes"],
            "eval_proactive_wire_bytes": source["eval_proactive_wire_bytes"],
            "full_reactive_wire_bytes": source["full_reactive_wire_bytes"],
            "full_proactive_wire_bytes": source["full_proactive_wire_bytes"],
            "transfer_count_evaluation": source["transfer_count"],
            "transfer_count_full_run": source["full_transfer_count"],
            "unused_transfer_wire_ratio": source["waste_ratio"],
            "waste_ratio_original_name": "waste_ratio" if source["waste_ratio"] is not None else None,
            "proactive_actions_evaluation": source["proactive_actions"],
            "proactive_actions_full_run": source["full_proactive_actions"],
            "reactive_actions_evaluation": source["reactive_actions"],
            "reactive_actions_full_run": source["full_reactive_actions"],
            "cache_eviction_count_all_causes_evaluation": source["cache_eviction_count_evaluation_all_causes"],
            "cache_eviction_count_all_causes_full_run": source["cache_eviction_count_full_run_all_causes"],
            "cache_admission_skips_evaluation": source["cache_admission_skip_evaluation"],
            "cache_admission_skips_full_run": source["cache_admission_skip_full_run"],
            "no_capacity_candidate_skips_evaluation": source["no_capacity_candidate_skips_evaluation"],
            "no_capacity_candidate_skips_full_run": source["no_capacity_candidate_skips_full_run"],
            "chain_depth_mean": source["chain_depth_mean"], "chain_depth_median": source["chain_depth_median"],
            "chain_depth_p90": source["chain_depth_p90"],
            "new_page_observed_use_ratio": source["new_page_observed_use_ratio"],
            "positive_candidate_count_evaluation": source["positive_candidate_count_evaluation"],
            "shortlist_count_evaluation": source["shortlist_count_evaluation"],
            "capacity_metric_scope": "D1_ALL_CAUSES_ONLY" if source["cache_eviction_count_evaluation_all_causes"] is not None else "NA_NOT_EMITTED",
            "notes": "Stage D1 fixed sensitivity case",
        }
        output.append(normalize_row(row))
    return output, source_rows


def d2_rows(gate: Mapping[str, Any], d2_table: Sequence[Mapping[str, Any]]) -> list[dict[str, Any]]:
    manifest_cases = {(c["workload"], c["case_id"]): c for c in gate["d2_manifest"]["cases"]}
    output: list[dict[str, Any]] = []
    for source in d2_table:
        if source["action"] != "MOVE":
            continue
        case = manifest_cases[(source["workload"], source["case_id"])]
        config = case["config"]
        trigger = source["trigger"]
        row = {
            "workload": source["workload"], "family": "MOVE", "referenced_by_families": "MOVE",
            "source_kind": "STAGE_D2_REPLAY", "source_run_id": source["source_run_id"],
            "source_artifact": "results/task_main/stage_d2_move_ablation/run_20260918_02/tables/copy_vs_move.json",
            "protocol_version": source["protocol_version"], "trace_sha256": config["trace_sha256"],
            "validation_status": "PASS", "deterministic": case["deterministic"], "line_id": None,
            "case_id": source["case_id"], "routing": "R_AFF", "policy": trigger,
            "action": "MOVE", "N": source["N"],
            "capacity_pages_per_pod": source["capacity_pages_per_pod"], "theta": config["theta"],
            "W_seconds": source["W_seconds"], "h_seconds": source["h_seconds"], "K": source["K"],
            "q": config["recency_quantile"], "evaluation_start_ms": source["evaluation_start_ms"],
            "evaluation_end_ms": source["evaluation_end_ms"], "saved_tokens": source["saved_tokens"],
            "weighted_saved_tokens": source["weighted_saved_tokens"],
            "request_hit_rate": source["request_hit_rate"], "token_hit_rate": source["token_hit_rate"],
            "skew_ratio": source["skew_ratio"], "eval_wire_bytes": source["eval_wire_bytes"],
            "full_wire_bytes": source["full_wire_bytes"], "eval_reactive_wire_bytes": 0,
            "eval_proactive_wire_bytes": source["eval_wire_bytes"], "full_reactive_wire_bytes": 0,
            "full_proactive_wire_bytes": source["full_wire_bytes"],
            "transfer_count_evaluation": source["proactive_action_count_evaluation"],
            "transfer_count_full_run": source["proactive_action_count_full_run"],
            "unused_transfer_wire_ratio": source["unused_transfer_wire_ratio"],
            "waste_ratio_original_name": "unused_move_wire_ratio",
            "proactive_actions_evaluation": source["proactive_action_count_evaluation"],
            "proactive_actions_full_run": source["proactive_action_count_full_run"],
            "reactive_actions_evaluation": 0, "reactive_actions_full_run": 0,
            "ordinary_lru_evictions_evaluation": source["ordinary_lru_evictions_evaluation"],
            "ordinary_lru_evictions_full_run": source["ordinary_lru_evictions_full_run"],
            "ordinary_cache_turnover_pages_evaluation": source["ordinary_cache_turnover_pages_evaluation"],
            "ordinary_cache_turnover_pages_full_run": source["ordinary_cache_turnover_pages_full_run"],
            "cache_admission_skips_evaluation": source["cache_admission_skips_evaluation"],
            "cache_admission_skips_full_run": source["cache_admission_skips_full_run"],
            "no_capacity_candidate_skips_evaluation": source["no_capacity_candidate_skips_evaluation"],
            "no_capacity_candidate_skips_full_run": source["no_capacity_candidate_skips_full_run"],
            "move_actions": source["move_actions"], "source_released_pages": source["source_released_pages"],
            "source_released_bytes": source["source_released_bytes"],
            "release_fraction_mean": source["release_fraction_mean"],
            "release_fraction_median": source["release_fraction_median"],
            "zero_release_actions": source["zero_release_actions"],
            "partial_release_actions": source["partial_release_actions"],
            "full_release_actions": source["full_release_actions"],
            "chain_depth_mean": source["chain_depth_mean"], "chain_depth_median": source["chain_depth_median"],
            "chain_depth_p90": source["chain_depth_p90"],
            "capacity_metric_scope": "D2_ORDINARY_LRU_EXCLUDES_MOVE_RELEASE",
            "notes": "Stage D2 COPY-then-safe-release MOVE ablation",
        }
        output.append(normalize_row(row))
    return output


def attach_reference_families(canonical: list[dict[str, Any]], d1_source: Sequence[Mapping[str, Any]],
                              d2_table: Sequence[Mapping[str, Any]]) -> None:
    by_key = {behavior_key(row): row for row in canonical}
    require(len(by_key) == len(canonical), "Canonical behavior key is duplicated before references")
    for source in d1_source:
        if source["source_kind"] != "FORMAL_REFERENCE":
            continue
        probe = {
            "workload": source["workload"], "routing": source["routing"], "policy": source["policy"],
            "action": source["action"], "N": source["N"], "theta": source["theta"],
            "W_seconds": source["W"], "h_seconds": source["h"], "K": source["K"], "q": source["q"],
            "evaluation_start_ms": source["evaluation_start"], "evaluation_end_ms": source["evaluation_end"],
        }
        target = by_key.get(behavior_key(probe))
        require(target is not None and target["source_kind"] == "FORMAL_MAIN",
                f"D1 formal reference does not resolve canonically: {source['workload']} {source['case_id']}")
        for metric, canonical_name in (("saved_tokens", "saved_tokens"),
                                       ("weighted_saved_tokens", "weighted_saved_tokens"),
                                       ("token_hit_rate", "token_hit_rate"), ("skew_ratio", "skew_ratio"),
                                       ("eval_wire_bytes", "eval_wire_bytes"), ("full_wire_bytes", "full_wire_bytes")):
            require(equal_number(source[metric], target[canonical_name]),
                    f"D1 formal reference value mismatch: {source['case_id']} {metric}")
        families = set(target["referenced_by_families"].split(";"))
        families.add(source["family"])
        target["referenced_by_families"] = ";".join(sorted(families))
    for source in d2_table:
        if source["action"] != "COPY":
            continue
        policy = source["trigger"]
        probe = {
            "workload": source["workload"], "routing": "R_AFF", "policy": policy, "action": "COPY",
            "N": source["N"], "theta": 2, "W_seconds": source["W_seconds"],
            "h_seconds": source["h_seconds"], "K": source["K"], "q": 0.9,
            "evaluation_start_ms": source["evaluation_start_ms"], "evaluation_end_ms": source["evaluation_end_ms"],
        }
        target = by_key.get(behavior_key(probe))
        require(target is not None and target["source_kind"] == "FORMAL_MAIN", "D2 COPY reference not canonical")
        for metric in ("saved_tokens", "weighted_saved_tokens", "token_hit_rate", "skew_ratio",
                       "eval_wire_bytes", "full_wire_bytes"):
            require(equal_number(source[metric], target[metric]), f"D2 COPY reference mismatch: {metric}")
        families = set(target["referenced_by_families"].split(";")); families.add("MOVE")
        target["referenced_by_families"] = ";".join(sorted(families))


def main_tables(canonical: Sequence[Mapping[str, Any]]) -> tuple[list[dict[str, Any]], ...]:
    main = sorted((r for r in canonical if r["family"] == "MAIN"), key=lambda r: (r["workload"], r["line_id"]))
    result: list[dict[str, Any]] = []
    for workload in ("conversation", "toolagent"):
        rows = [r for r in main if r["workload"] == workload]
        baseline = next(r for r in rows if r["line_id"] == 2)
        for r in rows:
            result.append({
                "workload": workload, "Line": r["line_id"], "Strategy": LINE_STRATEGY[r["line_id"]],
                "Saved_M": r["saved_tokens_million"],
                "Delta_Saved_vs_Line2_pct": percent_delta(r["saved_tokens"], baseline["saved_tokens"]),
                "Token_Hit_pct": r["token_hit_rate"] * 100,
                "WeightedSaved_M": r["weighted_saved_tokens_million"],
                "Delta_Weighted_vs_Line2_pct": percent_delta(r["weighted_saved_tokens"], baseline["weighted_saved_tokens"]),
                "Skew": r["skew_ratio"], "Eval_Wire_TB": r["eval_wire_tb_decimal"],
                "Full_Wire_TB": r["full_wire_tb_decimal"],
                "Unused_Wire_pct": None if r["unused_transfer_wire_ratio"] is None else r["unused_transfer_wire_ratio"] * 100,
                "canonical_case_id": r["canonical_case_id"], "baseline": "Line2 R_REQ_KV",
            })
    conversation = [r for r in result if r["workload"] == "conversation"]
    toolagent = [r for r in result if r["workload"] == "toolagent"]
    return conversation, toolagent, result


def baseline_table(canonical: Sequence[Mapping[str, Any]]) -> list[dict[str, Any]]:
    result = []
    for workload in ("conversation", "toolagent"):
        aff = next(r for r in canonical if r["workload"] == workload and r["family"] == "MAIN" and r["line_id"] == 1)
        req = next(r for r in canonical if r["workload"] == workload and r["family"] == "MAIN" and r["line_id"] == 2)
        least = next(r for r in canonical if r["workload"] == workload and r["family"] == "R_LEAST")
        for label, r in (("R_AFF", aff), ("R_LEAST", least), ("R_REQ_KV", req)):
            result.append({
                "workload": workload, "baseline": label, "Saved_M": r["saved_tokens_million"],
                "WeightedSaved_M": r["weighted_saved_tokens_million"], "TokenHit": r["token_hit_rate"],
                "Skew": r["skew_ratio"], "EvalWireTB": r["eval_wire_tb_decimal"],
                "FullWireTB": r["full_wire_tb_decimal"],
                "Delta_Saved_vs_R_AFF": r["saved_tokens"] - aff["saved_tokens"],
                "Delta_Saved_vs_R_AFF_pct": percent_delta(r["saved_tokens"], aff["saved_tokens"]),
                "Delta_Skew_vs_R_AFF": r["skew_ratio"] - aff["skew_ratio"],
                "main_task_baseline": label == "R_REQ_KV", "canonical_case_id": r["canonical_case_id"],
            })
    return result


def oracle_headroom(canonical: Sequence[Mapping[str, Any]]) -> list[dict[str, Any]]:
    result = []
    for workload in ("conversation", "toolagent"):
        baseline = next(r for r in canonical if r["workload"] == workload and r["family"] == "MAIN" and r["line_id"] == 2)
        oracle = next(r for r in canonical if r["workload"] == workload and r["family"] == "MAIN" and r["line_id"] == 3)
        result.append({
            "workload": workload, "reference_name": "Future-Demand Oracle Reference Headroom",
            "R_REQ_KV_Saved": baseline["saved_tokens"], "Oracle_Saved": oracle["saved_tokens"],
            "Delta_Saved": oracle["saved_tokens"] - baseline["saved_tokens"],
            "Delta_Saved_pct": percent_delta(oracle["saved_tokens"], baseline["saved_tokens"]),
            "R_REQ_KV_Weighted": baseline["weighted_saved_tokens"],
            "Oracle_Weighted": oracle["weighted_saved_tokens"],
            "Delta_Weighted": oracle["weighted_saved_tokens"] - baseline["weighted_saved_tokens"],
            "Delta_Weighted_pct": percent_delta(oracle["weighted_saved_tokens"], baseline["weighted_saved_tokens"]),
            "R_REQ_KV_Skew": baseline["skew_ratio"], "Oracle_Skew": oracle["skew_ratio"],
            "Delta_Skew": oracle["skew_ratio"] - baseline["skew_ratio"],
            "Oracle_Eval_Wire": oracle["eval_wire_bytes"], "Oracle_Full_Wire": oracle["full_wire_bytes"],
            "Oracle_Unused_Wire": oracle["unused_transfer_wire_ratio"],
            "comparison_status": "SAME_COHORT_COMPARABLE",
        })
    return result


def history_summary(canonical: Sequence[Mapping[str, Any]]) -> list[dict[str, Any]]:
    output = []
    for workload in ("conversation", "toolagent"):
        baseline = next(r for r in canonical if r["workload"] == workload and r["family"] == "MAIN" and r["line_id"] == 2)
        for line in (4, 5, 6):
            row = next(r for r in canonical if r["workload"] == workload and r["family"] == "MAIN" and r["line_id"] == line)
            output.append({
                "workload": workload, "line": line, "strategy": LINE_STRATEGY[line],
                "Delta_Saved_vs_R_REQ_KV": row["saved_tokens"] - baseline["saved_tokens"],
                "Delta_Saved_vs_R_REQ_KV_pct": percent_delta(row["saved_tokens"], baseline["saved_tokens"]),
                "Delta_Weighted_vs_R_REQ_KV": row["weighted_saved_tokens"] - baseline["weighted_saved_tokens"],
                "Delta_Weighted_vs_R_REQ_KV_pct": percent_delta(row["weighted_saved_tokens"], baseline["weighted_saved_tokens"]),
                "Skew": row["skew_ratio"], "Eval_Wire": row["eval_wire_bytes"],
                "Full_Wire": row["full_wire_bytes"], "Waste": row["unused_transfer_wire_ratio"],
                "canonical_case_id": row["canonical_case_id"],
            })
    return output


def sensitivity_context(canonical: Sequence[Mapping[str, Any]], d1_source: Sequence[Mapping[str, Any]]) -> list[dict[str, Any]]:
    by_key = {behavior_key(r): r for r in canonical}
    output = []
    for source in d1_source:
        if source["family"] not in {"THETA", "PERSISTENCE_K", "RECENCY_Q", "ORACLE_W", "N"}:
            continue
        probe = {
            "workload": source["workload"], "routing": source["routing"], "policy": source["policy"],
            "action": source["action"], "N": source["N"], "theta": source["theta"],
            "W_seconds": source["W"], "h_seconds": source["h"], "K": source["K"], "q": source["q"],
            "evaluation_start_ms": source["evaluation_start"], "evaluation_end_ms": source["evaluation_end"],
        }
        canonical_row = by_key.get(behavior_key(probe))
        require(canonical_row is not None, f"Sensitivity context cannot resolve {source['case_id']}")
        output.append({
            "canonical_case_id": canonical_row["canonical_case_id"], "workload": source["workload"],
            "family": source["family"], "case_id": source["case_id"], "source_kind": source["source_kind"],
            "source_run_id": source["source_run_id"], "routing": source["routing"], "policy": source["policy"],
            "action": source["action"], "N": source["N"],
            "capacity_pages_per_pod": source["capacity_pages_per_pod"],
            "total_cluster_capacity_pages": source["total_cluster_capacity_pages"], "theta": source["theta"],
            "W_seconds": source["W"], "h_seconds": source["h"], "K": source["K"], "q": source["q"],
            "evaluation_start_ms": source["evaluation_start"], "evaluation_end_ms": source["evaluation_end"],
            "evaluation_label": cohort_label(source["evaluation_start"], source["evaluation_end"]),
            "saved_tokens": canonical_row["saved_tokens"], "weighted_saved_tokens": canonical_row["weighted_saved_tokens"],
            "request_hit_rate": canonical_row["request_hit_rate"], "token_hit_rate": canonical_row["token_hit_rate"],
            "skew_ratio": canonical_row["skew_ratio"], "eval_wire_bytes": canonical_row["eval_wire_bytes"],
            "full_wire_bytes": canonical_row["full_wire_bytes"],
            "eval_reactive_wire_bytes": canonical_row["eval_reactive_wire_bytes"],
            "eval_proactive_wire_bytes": canonical_row["eval_proactive_wire_bytes"],
            "full_reactive_wire_bytes": canonical_row["full_reactive_wire_bytes"],
            "full_proactive_wire_bytes": canonical_row["full_proactive_wire_bytes"],
            "waste_ratio": canonical_row["unused_transfer_wire_ratio"],
            "proactive_actions_evaluation": source["proactive_actions"],
            "proactive_actions_full_run": source["full_proactive_actions"],
            "reactive_actions_evaluation": source["reactive_actions"],
            "reactive_actions_full_run": source["full_reactive_actions"],
            "gate_pass_count_evaluation": source["gate_pass_count_evaluation"],
            "reactive_copy_started_evaluation": source["reactive_copy_started_evaluation"],
            "chain_depth_mean": source["chain_depth_mean"], "chain_depth_median": source["chain_depth_median"],
            "chain_depth_p90": source["chain_depth_p90"],
            "new_page_observed_use_ratio": source["new_page_observed_use_ratio"],
            "positive_candidate_count_evaluation": source["positive_candidate_count_evaluation"],
            "shortlist_count_evaluation": source["shortlist_count_evaluation"],
            "comparison_reference_case_id": source["comparison_reference_case_id"],
        })
    return sorted(output, key=lambda r: (r["family"], r["workload"], r["h_seconds"], r["K"], r["q"], r["W_seconds"], r["N"], r["theta"]))


def sensitivity_subtables(rows: Sequence[Mapping[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    theta, persistence, recency, oracle, n_rows = [], [], [], [], []
    for r in rows:
        if r["family"] == "THETA":
            theta.append({
                "workload": r["workload"], "theta": r["theta"], "Saved": r["saved_tokens"],
                "Weighted": r["weighted_saved_tokens"], "Skew": r["skew_ratio"],
                "Full_reactive_wire": r["full_reactive_wire_bytes"],
                "gate_count": r["gate_pass_count_evaluation"],
                "reactive_copy_count": r["reactive_copy_started_evaluation"],
                "evaluation_label": r["evaluation_label"], "canonical_case_id": r["canonical_case_id"],
            })
        elif r["family"] == "PERSISTENCE_K":
            persistence.append({
                "workload": r["workload"], "h": r["h_seconds"], "K": r["K"],
                "Saved": r["saved_tokens"], "Weighted": r["weighted_saved_tokens"], "Skew": r["skew_ratio"],
                "Eval_Wire": r["eval_wire_bytes"], "Full_Wire": r["full_wire_bytes"],
                "Waste": r["waste_ratio"], "proactive_actions": r["proactive_actions_evaluation"],
                "chain_depth_mean": r["chain_depth_mean"],
                "new_page_use_ratio": r["new_page_observed_use_ratio"],
                "evaluation_label": r["evaluation_label"], "canonical_case_id": r["canonical_case_id"],
            })
        elif r["family"] == "RECENCY_Q":
            recency.append({
                "workload": r["workload"], "q": r["q"], "Saved": r["saved_tokens"],
                "Weighted": r["weighted_saved_tokens"], "Skew": r["skew_ratio"],
                "Wire": r["eval_wire_bytes"], "Full_Wire": r["full_wire_bytes"], "Waste": r["waste_ratio"],
                "positive_candidates": r["positive_candidate_count_evaluation"],
                "shortlist_count": r["shortlist_count_evaluation"],
                "actions": r["proactive_actions_evaluation"], "evaluation_label": r["evaluation_label"],
                "canonical_case_id": r["canonical_case_id"],
            })
        elif r["family"] == "ORACLE_W":
            require(r["evaluation_label"] == "W_SENS_10_25", "Oracle W row is outside common-support cohort")
            oracle.append({
                "workload": r["workload"], "W_minutes": r["W_seconds"] / 60,
                "Saved": r["saved_tokens"], "Weighted": r["weighted_saved_tokens"], "Skew": r["skew_ratio"],
                "Eval_Wire": r["eval_wire_bytes"], "Full_Wire": r["full_wire_bytes"],
                "Unused_ratio": r["waste_ratio"], "actions": r["proactive_actions_evaluation"],
                "chain_depth": r["chain_depth_mean"], "evaluation_label": r["evaluation_label"],
                "canonical_case_id": r["canonical_case_id"],
            })
        elif r["family"] == "N":
            n_rows.append({
                "workload": r["workload"], "N": r["N"],
                "total_cluster_capacity_pages": r["total_cluster_capacity_pages"],
                "strategy": f"{r['routing']} / {r['policy']}", "Saved": r["saved_tokens"],
                "Weighted": r["weighted_saved_tokens"], "Skew": r["skew_ratio"],
                "Wire": r["eval_wire_bytes"], "Full_Wire": r["full_wire_bytes"], "Waste": r["waste_ratio"],
                "evaluation_label": r["evaluation_label"],
                "warning": "per-Pod capacity fixed; total cluster capacity changes with N.",
                "canonical_case_id": r["canonical_case_id"],
            })
    return {"theta_sensitivity_final": theta, "persistence_sensitivity_final": persistence,
            "recency_sensitivity_final": recency, "oracle_w_sensitivity_final": oracle,
            "n_sensitivity_final": n_rows}


def copy_move_tables(d2_table: Sequence[Mapping[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    rows = [dict(r) for r in d2_table]
    grouped: dict[tuple[str, str], dict[str, Mapping[str, Any]]] = defaultdict(dict)
    for row in rows:
        grouped[(row["workload"], row["trigger"])][row["action"]] = row
    deltas = []
    for (workload, trigger), pair in sorted(grouped.items()):
        require(set(pair) == {"COPY", "MOVE"}, f"Incomplete COPY/MOVE pair: {workload} {trigger}")
        copy, move = pair["COPY"], pair["MOVE"]
        deltas.append({
            "workload": workload, "trigger": trigger, "copy_case_id": copy["case_id"],
            "move_case_id": move["case_id"], "delta_saved_tokens": move["saved_tokens"] - copy["saved_tokens"],
            "delta_saved_pct": percent_delta(move["saved_tokens"], copy["saved_tokens"]),
            "delta_weighted": move["weighted_saved_tokens"] - copy["weighted_saved_tokens"],
            "delta_weighted_pct": percent_delta(move["weighted_saved_tokens"], copy["weighted_saved_tokens"]),
            "delta_token_hit_pp": (move["token_hit_rate"] - copy["token_hit_rate"]) * 100,
            "delta_skew": move["skew_ratio"] - copy["skew_ratio"],
            "delta_eval_wire_bytes": move["eval_wire_bytes"] - copy["eval_wire_bytes"],
            "delta_eval_wire_pct": percent_delta(move["eval_wire_bytes"], copy["eval_wire_bytes"]),
            "delta_full_wire_bytes": move["full_wire_bytes"] - copy["full_wire_bytes"],
            "delta_full_wire_pct": percent_delta(move["full_wire_bytes"], copy["full_wire_bytes"]),
            "delta_unused_wire_pp": (move["unused_transfer_wire_ratio"] - copy["unused_transfer_wire_ratio"]) * 100,
            "delta_evictions": move["ordinary_lru_evictions_evaluation"] - copy["ordinary_lru_evictions_evaluation"],
            "delta_evictions_pct": percent_delta(move["ordinary_lru_evictions_evaluation"], copy["ordinary_lru_evictions_evaluation"]),
            "released_pages": move["source_released_pages"], "release_fraction": move["release_fraction_mean"],
            "zero_release_actions": move["zero_release_actions"],
            "partial_release_actions": move["partial_release_actions"],
            "full_release_actions": move["full_release_actions"],
            "comparison_status": "SAME_COHORT_COMPARABLE",
        })
    return rows, deltas


def prompt_bucket_tables() -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    source = load_json(FORMAL / "tables/prompt_bucket.json")
    by_key = {(r["workload"], r["line_id"], r["bucket"]): r for r in source}
    rows = []
    oracle = []
    for r in source:
        baseline = by_key[(r["workload"], 2, r["bucket"])]
        row = {
            "workload": r["workload"], "line": r["line_id"], "strategy": LINE_STRATEGY[r["line_id"]],
            "bucket": r["bucket"], "request_count": r["request_count"], "input_tokens": r["input_tokens"],
            "saved_tokens": r["saved_tokens"], "token_hit_rate": r["token_hit_rate"],
            "weighted_saved_tokens": r["weighted_saved_tokens"],
            "delta_saved_vs_line2": r["saved_tokens"] - baseline["saved_tokens"],
            "delta_weighted_vs_line2": r["weighted_saved_tokens"] - baseline["weighted_saved_tokens"],
            "evaluation_label": "MAIN_25_45",
        }
        rows.append(row)
        if r["line_id"] == 3:
            oracle.append(dict(row))
    return rows, oracle


def diagnostics_table() -> list[dict[str, Any]]:
    efficiency_rows = load_json(DIAGNOSTICS / "strategy_copy_efficiency.json")
    efficiency = {(r["workload"], r["line_id"]): r for r in efficiency_rows}
    time_rows = load_json(DIAGNOSTICS / "time_to_use_summary.json")
    times = {(r["workload"], r["line_id"]): r for r in time_rows}
    depth_rows = [r for r in load_json(DIAGNOSTICS / "copy_depth_summary.json") if r["group"] == "ALL"]
    depths = {(r["workload"], r["line_id"]): r for r in depth_rows}
    churn_rows = load_json(DIAGNOSTICS / "cache_churn_summary.json")
    churn = {(r["workload"], r["line_id"]): r for r in churn_rows}
    raw = load_json(DIAGNOSTICS / "copy_action_diagnostics.json")["actions"]
    fractions: dict[tuple[str, int], list[float]] = defaultdict(list)
    for action in raw:
        value = action["observed_chain_residency_reuse_fraction"]
        if value is not None:
            fractions[(action["workload"], action["line_id"])].append(value)
    output = []
    for workload in ("conversation", "toolagent"):
        for line in range(1, 8):
            if line <= 2:
                output.append({
                    "workload": workload, "line": line, "strategy": LINE_STRATEGY[line],
                    "full_chain_unused_ratio": None, "unused_transfer_wire_ratio": None,
                    "partial_reuse_any_ratio": None, "partial_reuse_fraction_mean": None,
                    "partial_reuse_fraction_median": None, "new_page_observed_use_ratio": None,
                    "wire_redundancy_ratio": None, "chain_depth_mean": None,
                    "chain_depth_median": None, "chain_depth_p90": None,
                    "ordinary_eviction_count": None, "cache_turnover": None,
                    "time_to_full_use_p50_ms": None, "time_to_full_use_p90_ms": None,
                    "diagnostic_na_reason": "NO_PROACTIVE_ACTIONS",
                })
                continue
            e, t, d = efficiency[(workload, line)], times[(workload, line)], depths[(workload, line)]
            c = churn.get((workload, line))
            output.append({
                "workload": workload, "line": line, "strategy": LINE_STRATEGY[line],
                "full_chain_unused_ratio": e["full_chain_unused_action_count"] / e["evaluation_action_count"],
                "unused_transfer_wire_ratio": e["formal_waste_ratio"],
                "partial_reuse_any_ratio": t["any_prefix_reused_action_count"] / e["evaluation_action_count"],
                "partial_reuse_fraction_mean": math.fsum(fractions[(workload, line)]) / len(fractions[(workload, line)]),
                "partial_reuse_fraction_median": d["partial_reuse_fraction_median"],
                "new_page_observed_use_ratio": e["new_page_observed_use_ratio"],
                "wire_redundancy_ratio": e["wire_redundancy_token_weighted_overall_ratio"],
                "chain_depth_mean": e["chain_depth_pages_mean"],
                "chain_depth_median": e["chain_depth_pages_median"], "chain_depth_p90": e["chain_depth_pages_p90"],
                "ordinary_eviction_count": None if c is None else c["evaluation_cache_eviction_records_all_causes"],
                "cache_turnover": None if c is None else c["cache_turnover_pages_all_causes"],
                "time_to_full_use_p50_ms": t["full_chain_time_to_use_ms_median"],
                "time_to_full_use_p90_ms": t["full_chain_time_to_use_ms_p90"],
                "diagnostic_na_reason": None if c is not None else "CHURN_NOT_EMITTED_FOR_THIS_LINE_BY_DIAGNOSTICS_V1",
            })
    return output


def strategy_master(main_rows: Sequence[Mapping[str, Any]], baseline: Sequence[Mapping[str, Any]],
                    sensitivity: Sequence[Mapping[str, Any]], copy_move: Sequence[Mapping[str, Any]],
                    canonical: Sequence[Mapping[str, Any]]) -> list[dict[str, Any]]:
    by_id = {r["canonical_case_id"]: r for r in canonical}
    contexts: list[dict[str, Any]] = []
    for r in main_rows:
        c = by_id[r["canonical_case_id"]]
        contexts.append({"canonical": c, "family": "MAIN", "strategy": r["Strategy"],
                         "reference_label": "Line2 R_REQ_KV", "reference_selector": (c["workload"], "MAIN_LINE2")})
    for r in baseline:
        if r["baseline"] != "R_LEAST":
            continue
        c = by_id[r["canonical_case_id"]]
        contexts.append({"canonical": c, "family": "R_LEAST", "strategy": "R_LEAST",
                         "reference_label": "R_AFF", "reference_selector": (c["workload"], "R_AFF")})
    for r in sensitivity:
        c = by_id[r["canonical_case_id"]]
        contexts.append({"canonical": c, "family": r["family"],
                         "strategy": f"{r['routing']} / {r['policy']}",
                         "reference_label": r["comparison_reference_case_id"],
                         "reference_context": r})
    copy_by = {(r["workload"], r["trigger"]): r for r in copy_move if r["action"] == "COPY"}
    canonical_key = {behavior_key(r): r for r in canonical}
    for r in copy_move:
        trigger = r["trigger"]
        probe = {"workload": r["workload"], "routing": "R_AFF", "policy": trigger,
                 "action": r["action"], "N": r["N"], "theta": 2, "W_seconds": r["W_seconds"],
                 "h_seconds": r["h_seconds"], "K": r["K"], "q": .9,
                 "evaluation_start_ms": r["evaluation_start_ms"], "evaluation_end_ms": r["evaluation_end_ms"]}
        c = canonical_key[behavior_key(probe)]
        contexts.append({"canonical": c, "family": "MOVE", "strategy": f"{trigger} {r['action']}",
                         "reference_label": f"{trigger} COPY", "copy_reference": copy_by[(r["workload"], trigger)]})

    main_line2 = {(r["workload"]): by_id[r["canonical_case_id"]] for r in main_rows if r["Line"] == 2}
    aff = {(r["workload"]): by_id[r["canonical_case_id"]] for r in main_rows if r["Line"] == 1}
    output = []
    for context in contexts:
        c, family = context["canonical"], context["family"]
        if family == "MAIN":
            ref = main_line2[c["workload"]]
        elif family == "R_LEAST":
            ref = aff[c["workload"]]
        elif family == "MOVE":
            cp = context["copy_reference"]
            ref_probe = {"workload": cp["workload"], "routing": "R_AFF", "policy": cp["trigger"],
                         "action": "COPY", "N": cp["N"], "theta": 2, "W_seconds": cp["W_seconds"],
                         "h_seconds": cp["h_seconds"], "K": cp["K"], "q": .9,
                         "evaluation_start_ms": cp["evaluation_start_ms"], "evaluation_end_ms": cp["evaluation_end_ms"]}
            ref = canonical_key[behavior_key(ref_probe)]
        else:
            rc = context["reference_context"]
            candidates = [x for x in sensitivity if x["workload"] == c["workload"] and x["family"] == family]
            if family == "THETA": candidates = [x for x in candidates if float(x["theta"]) == 2.0]
            elif family == "PERSISTENCE_K": candidates = [x for x in candidates if x["h_seconds"] == rc["h_seconds"] and x["K"] == 10]
            elif family == "RECENCY_Q": candidates = [x for x in candidates if float(x["q"]) == .9]
            elif family == "ORACLE_W": candidates = [x for x in candidates if float(x["W_seconds"]) == 300.0]
            elif family == "N": candidates = [x for x in candidates if x["N"] == 4 and x["routing"] == rc["routing"] and x["policy"] == rc["policy"]]
            require(len(candidates) == 1, f"Sensitivity reference is ambiguous: {family} {c['canonical_case_id']}")
            ref = by_id[candidates[0]["canonical_case_id"]]
        same_cohort = (c["workload"] == ref["workload"] and c["evaluation_start_ms"] == ref["evaluation_start_ms"]
                       and c["evaluation_end_ms"] == ref["evaluation_end_ms"])
        output.append({
            "Workload": c["workload"], "Family": family, "Strategy": context["strategy"],
            "Routing": c["routing"], "Action": c["action"], "N": c["N"], "theta": c["theta"],
            "W_min": c["W_seconds"] / 60, "h_min": c["h_seconds"] / 60, "K": c["K"], "q": c["q"],
            "Eval_Cohort": c["evaluation_label"], "Saved_M": c["saved_tokens_million"],
            "WeightedSaved_M": c["weighted_saved_tokens_million"], "TokenHit_pct": c["token_hit_rate"] * 100,
            "Skew": c["skew_ratio"], "EvalWire_TB": c["eval_wire_tb_decimal"],
            "FullWire_TB": c["full_wire_tb_decimal"],
            "UnusedWire_pct": None if c["unused_transfer_wire_ratio"] is None else c["unused_transfer_wire_ratio"] * 100,
            "DeltaSaved_vs_Reference_pct": percent_delta(c["saved_tokens"], ref["saved_tokens"]) if same_cohort else None,
            "Reference": context["reference_label"],
            "comparison_status": "SAME_COHORT_COMPARABLE" if same_cohort else "CROSS_COHORT_NOT_COMPARABLE",
            "canonical_case_id": c["canonical_case_id"],
            "Notes": "WeightedSaved is a multiplier-weighted token benefit proxy; wire uses decimal TB.",
        })
    return output


def validate_all(canonical: Sequence[Mapping[str, Any]], main_rows: Sequence[Mapping[str, Any]],
                 sensitivity: Sequence[Mapping[str, Any]], copy_move: Sequence[Mapping[str, Any]],
                 copy_delta: Sequence[Mapping[str, Any]], buckets: Sequence[Mapping[str, Any]]) -> dict[str, bool]:
    checks: dict[str, bool] = {}
    checks["formal_main_14_rows"] = sum(r["family"] == "MAIN" for r in canonical) == 14
    checks["formal_7_per_workload"] = all(sum(r["family"] == "MAIN" and r["workload"] == w for r in canonical) == 7 for w in TRACE_SHA256)
    checks["line5_line7_main_metrics_equal"] = all(
        all(equal_number(next(r for r in canonical if r["family"] == "MAIN" and r["workload"] == w and r["line_id"] == 5)[m],
                         next(r for r in canonical if r["family"] == "MAIN" and r["workload"] == w and r["line_id"] == 7)[m])
            for m in ("saved_tokens", "weighted_saved_tokens", "request_hit_rate", "token_hit_rate", "skew_ratio",
                      "eval_wire_bytes", "full_wire_bytes", "unused_transfer_wire_ratio")) for w in TRACE_SHA256)
    checks["rleast_both_workloads"] = {r["workload"] for r in canonical if r["family"] == "R_LEAST"} == set(TRACE_SHA256)
    checks["theta_grid"] = all({float(r["theta"]) for r in sensitivity if r["family"] == "THETA" and r["workload"] == w} == {1.5, 2.0, 3.0} for w in TRACE_SHA256)
    checks["persistence_grid"] = all({(float(r["h_seconds"]), r["K"]) for r in sensitivity if r["family"] == "PERSISTENCE_K" and r["workload"] == w} == {(60.0,k) for k in (5,10,20)} | {(300.0,k) for k in (5,10,20)} for w in TRACE_SHA256)
    checks["recency_grid"] = all({float(r["q"]) for r in sensitivity if r["family"] == "RECENCY_Q" and r["workload"] == w} == {.8,.9} for w in TRACE_SHA256)
    checks["oracle_w_grid_common_cohort"] = all({float(r["W_seconds"]) for r in sensitivity if r["family"] == "ORACLE_W" and r["workload"] == w} == {60.,300.,1800.} and all(r["evaluation_label"] == "W_SENS_10_25" for r in sensitivity if r["family"] == "ORACLE_W" and r["workload"] == w) for w in TRACE_SHA256)
    checks["n_grid"] = all({r["N"] for r in sensitivity if r["family"] == "N" and r["workload"] == w} == {2,4} for w in TRACE_SHA256)
    checks["copy_move_four_pairs"] = len(copy_delta) == 4 and len(copy_move) == 8
    bucket_groups: dict[tuple[str,int], list[Mapping[str,Any]]] = defaultdict(list)
    for row in buckets: bucket_groups[(row["workload"], row["line"])].append(row)
    checks["prompt_bucket_six_and_conserved"] = len(bucket_groups) == 14 and all(
        len(rows) == 6 and sum(r["saved_tokens"] for r in rows) == next(x for x in canonical if x["family"] == "MAIN" and x["workload"] == key[0] and x["line_id"] == key[1])["saved_tokens"]
        and math.isclose(math.fsum(r["weighted_saved_tokens"] for r in rows), next(x for x in canonical if x["family"] == "MAIN" and x["workload"] == key[0] and x["line_id"] == key[1])["weighted_saved_tokens"], rel_tol=1e-12, abs_tol=1e-7)
        for key, rows in bucket_groups.items())
    checks["wire_eval_le_full"] = all(r["eval_wire_bytes"] <= r["full_wire_bytes"] for r in canonical)
    checks["rates_in_unit_interval"] = all(v is None or 0 <= v <= 1 for r in canonical for v in (r["request_hit_rate"], r["token_hit_rate"]))
    checks["waste_in_unit_interval_or_na"] = all(r["unused_transfer_wire_ratio"] is None or 0 <= r["unused_transfer_wire_ratio"] <= 1 for r in canonical)
    moves = [r for r in canonical if r["action"] == "MOVE"]
    checks["move_release_partition"] = all(r["zero_release_actions"] + r["partial_release_actions"] + r["full_release_actions"] == r["move_actions"] for r in moves)
    checks["move_release_pages_nonnegative"] = all(r["source_released_pages"] >= 0 for r in moves)
    checks["canonical_id_unique"] = len({r["canonical_case_id"] for r in canonical}) == len(canonical)
    checks["behavior_key_unique"] = len({behavior_key(r) for r in canonical}) == len(canonical)
    checks["canonical_row_count_50"] = len(canonical) == 50
    checks["all_validation_pass"] = all(r["validation_status"] == "PASS" and r["deterministic"] for r in canonical)
    checks["all_protocol_frozen"] = all(r["protocol_version"] == PROTOCOL for r in canonical)
    checks["trace_identity_frozen"] = all(r["trace_sha256"] == TRACE_SHA256[r["workload"]] for r in canonical)
    checks["move_deltas_recomputed"] = all(
        d["comparison_status"] == "SAME_COHORT_COMPARABLE" and
        d["delta_saved_tokens"] == next(r for r in copy_move if r["workload"] == d["workload"] and r["trigger"] == d["trigger"] and r["action"] == "MOVE")["saved_tokens"] - next(r for r in copy_move if r["workload"] == d["workload"] and r["trigger"] == d["trigger"] and r["action"] == "COPY")["saved_tokens"]
        for d in copy_delta)
    checks["no_cross_cohort_delta"] = all(r["evaluation_label"] == "W_SENS_10_25" for r in sensitivity if r["family"] == "ORACLE_W")
    require(all(checks.values()), "Final consistency checks failed: " + ", ".join(k for k,v in checks.items() if not v))
    return checks


def data_dictionary(path: Path) -> None:
    definitions = {
        "canonical_case_id": ("Stable identity for one behavior configuration and Evaluation cohort.", "identifier", "No", "Never NA"),
        "referenced_by_families": ("Semicolon-separated contexts that reference the canonical row without duplicating it.", "labels", "No", "Never NA"),
        "saved_tokens": ("Sum of actual longest-Prefix hit tokens for requests in the Evaluation arrival cohort.", "tokens", "Only same cohort", "Never NA"),
        "saved_tokens_million": ("saved_tokens divided by 1e6.", "million tokens", "Only same cohort", "Never NA"),
        "weighted_saved_tokens": ("Multiplier-weighted Saved Tokens benefit proxy; it is not time or milliseconds.", "proxy units", "Only same cohort", "Never NA"),
        "weighted_saved_tokens_million": ("weighted_saved_tokens divided by 1e6.", "million proxy units", "Only same cohort", "Never NA"),
        "request_hit_rate": ("Fraction of Evaluation requests with a positive actual Prefix hit.", "0-1 rate", "Only same cohort", "NA only for an empty cohort"),
        "token_hit_rate": ("saved_tokens divided by Evaluation input tokens.", "0-1 rate", "Only same cohort", "NA only for zero input tokens"),
        "skew_ratio": ("Median per-window actual request-distribution Gini divided by uniform-random mean Gini.", "ratio", "Only same cohort/window definition/N", "NA if random mean is zero"),
        "eval_wire_bytes": ("Wire bytes for transfers whose start time is in Evaluation.", "bytes", "Only same cohort", "Never NA"),
        "eval_wire_tb_decimal": ("eval_wire_bytes divided by 1e12; decimal TB, not TiB.", "decimal TB", "Only same cohort", "Never NA"),
        "full_wire_bytes": ("Wire bytes for all transfers started in the full replay.", "bytes", "Compare only equal replay scope", "Never NA"),
        "full_wire_tb_decimal": ("full_wire_bytes divided by 1e12; decimal TB.", "decimal TB", "Compare only equal replay scope", "Never NA"),
        "unused_transfer_wire_ratio": ("Observed transferred-token fraction lacking full-chain generation-matched use within W; MOVE uses the same target-side rule.", "0-1 rate", "Only equal observation rule/W/cohort", "NA when no observed proactive transfer tokens"),
        "waste_ratio_original_name": ("Name of the source artifact field mapped to unused_transfer_wire_ratio.", "field name", "No", "NA when the ratio is NA"),
        "ordinary_lru_evictions_evaluation": ("Ordinary LRU eviction count during Evaluation; MOVE safe-release deletions are excluded.", "count", "Only equal scope", "NA if source emitted only all-cause or no churn diagnostics"),
        "cache_admission_skips_evaluation": ("Prompt Cache admission failures during Evaluation.", "count", "Only equal scope", "NA if not emitted"),
        "no_capacity_candidate_skips_evaluation": ("Proactive candidates skipped for target capacity during Evaluation.", "count", "Only equal scope", "NA if not emitted"),
        "move_actions": ("Evaluation proactive MOVE actions classified for source release.", "count", "Only MOVE", "NA for COPY"),
        "source_released_pages": ("Source pages safely released after target commit across Evaluation MOVE actions.", "pages", "Only MOVE", "NA for COPY"),
        "source_released_bytes": ("source_released_pages multiplied by 14 MiB/page.", "bytes", "Only MOVE", "NA for COPY"),
        "release_fraction_mean": ("Mean per-action safely released pages divided by moved-chain pages.", "0-1 fraction", "Only MOVE", "NA for COPY"),
        "release_fraction_median": ("Median per-action safe release fraction.", "0-1 fraction", "Only MOVE", "NA for COPY"),
        "zero_release_actions": ("MOVE actions that could release no source page.", "count", "Only MOVE", "NA for COPY"),
        "partial_release_actions": ("MOVE actions that released a strict suffix but not the full chain.", "count", "Only MOVE", "NA for COPY"),
        "full_release_actions": ("MOVE actions that released the full moved chain.", "count", "Only MOVE", "NA for COPY"),
        "evaluation_label": ("Named Evaluation cohort. MAIN_25_45 and W_SENS_10_25 are deliberately separate.", "label", "No", "Never NA"),
        "capacity_metric_scope": ("Provenance/scope of capacity and churn fields.", "label", "No", "Never NA"),
    }
    lines = ["# TaskMain Final Results Data Dictionary", "", "CSV uses `NA`; JSON uses `null`. Rates are stored on a 0–1 scale unless a field name ends in `_pct` or `_pp`.", "",
             "| Field | Definition | Unit | Cross-cohort comparison | NA meaning |", "|---|---|---|---|---|"]
    for field in FINAL_LONG_COLUMNS:
        definition = definitions.get(field, (field.replace("_", " ") + ".", "see field name", "Only identical cohort and scope", "Not applicable or not reliably emitted"))
        lines.append(f"| `{field}` | {definition[0]} | {definition[1]} | {definition[2]} | {definition[3]} |")
    lines += ["", "## Comparison rule", "", "Derived deltas and percentages are emitted only when workload, Evaluation start, and Evaluation end match. Otherwise `comparison_status=CROSS_COHORT_NOT_COMPARABLE` and the delta is NA.", "",
              "## Reporting-table derived fields", "",
              "- Fields ending in `_M` divide the corresponding raw token/proxy value by 1e6.",
              "- Fields ending in `_TB` use decimal TB (`bytes / 1e12`).",
              "- Fields ending in `_pct` are percentage-valued numbers; `_pp` is an absolute percentage-point delta.",
              "- `delta_*_pct` is `(case-reference)/reference*100`, emitted only for a same-workload, same-cohort comparison with a nonzero reference.",
              "- `full_chain_unused_ratio` in `diagnostics_summary` is action-weighted; `unused_transfer_wire_ratio` is transferred-token-weighted.",
              "- `partial_reuse_any_ratio` is the fraction of Evaluation actions with any generation-matched Prefix reuse.",
              "- `partial_reuse_fraction_mean/median` summarize the maximum generation-matched reused Prefix depth divided by chain depth.",
              "- `wire_redundancy_ratio` is the token-weighted share already resident on the target at ready; it is not causal benefit.",
              "- `time_to_full_use_p50_ms/p90_ms` is measured from ready to first complete generation-matched chain use.", "",
              "## Interpretation limits", "", "WeightedSaved is a token-benefit proxy. Skew is a ratio, not a percentage. Unused transfer wire requires full-chain use and does not imply that no partial Prefix was reused. MOVE source release is separate from target-side unused transfer wire.", ""]
    path.write_text("\n".join(lines), encoding="utf-8")


def build_xlsx(path: Path, sheets: Sequence[tuple[str, Sequence[Mapping[str, Any]]]]) -> None:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    fixed = datetime(2026, 9, 18, tzinfo=timezone.utc)
    wb.properties.created = fixed; wb.properties.modified = fixed
    readme = wb.create_sheet("README")
    readme_rows = [
        ("TaskMain Final Results", "Canonical reporting workbook; CSV/JSON remain machine-readable sources."),
        ("Status", "FINAL_RESULTS_CONSOLIDATION_PASS"),
        ("Cohorts", "MAIN_25_45=[1500000,2700000) ms; W_SENS_10_25=[600000,1500000) ms."),
        ("Saved", "Actual hit tokens; *_M divides by 1e6."),
        ("WeightedSaved", "Multiplier-weighted token benefit proxy, not milliseconds."),
        ("Wire", "Decimal TB = bytes/1e12; not TiB."),
        ("Rates", "Stored as 0-1 unless named _pct/_pp."),
        ("NA", "Not applicable, empty denominator, or not reliably emitted; see DATA_DICTIONARY.md."),
        ("Provenance", f"Formal {FORMAL_RUN_ID}; D1 run_20260917_01; D2 run_20260918_02; protocol {PROTOCOL}."),
        ("Assumptions", "N=4 and 585 pages/Pod unless a sensitivity row says otherwise; P-only; 512 tokens/page; 14 MiB/page; 25 GB/s per transfer."),
    ]
    for row in readme_rows: readme.append(row)
    readme.column_dimensions["A"].width = 22; readme.column_dimensions["B"].width = 110
    readme.freeze_panes = "A2"; readme["A1"].font = Font(bold=True); readme["B1"].font = Font(bold=True)

    for name, rows in sheets:
        ws = wb.create_sheet(name[:31])
        columns = list(rows[0]) if rows else []
        ws.append(columns)
        for row in rows:
            ws.append(["NA" if row.get(c) is None else row.get(c) for c in columns])
        ws.freeze_panes = "A2"
        if columns:
            ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for index, column in enumerate(columns, 1):
            values = [str(column)] + [str(r.get(column, "NA")) for r in rows[:1000]]
            ws.column_dimensions[get_column_letter(index)].width = min(max(len(v) for v in values) + 2, 50)
            lower = column.lower()
            if lower.endswith("_pct") or lower.endswith("_pp"):
                for cell in ws[get_column_letter(index)][1:]:
                    if isinstance(cell.value, (int, float)): cell.number_format = '0.0000"%"'
    wb.save(path)
    check = load_workbook(path, read_only=True, data_only=True)
    require(set(ws.title for ws in check.worksheets) >= {"README", "Main_7Line", "Baselines", "Oracle_Headroom", "Persistence", "Recency", "Oracle_W", "Theta", "N", "COPY_vs_MOVE", "Prompt_Buckets", "Diagnostics", "All_Long"}, "Excel sheets missing")
    check.close()


def audit_report(path: Path, output: Path, checks: Mapping[str, bool], row_counts: Mapping[str, int],
                 source_sha: Mapping[str, str], artifacts: Mapping[str, str]) -> None:
    lines = [
        "# TaskMain Final Results Consolidation", "", "**状态：FINAL_RESULTS_CONSOLIDATION_PASS。** 本阶段只读取既有 PASS artifacts；未调用 simulator，未修改既有结果、协议或配置。", "",
        "## 1. 数据来源", "",
        f"- Source A：Formal Main 14-line，run `{FORMAL_RUN_ID}`，manifest `{source_sha['formal_validation_manifest']}`。",
        "- Source B：Formal Result Diagnostics V1，仅用于诊断字段。",
        "- Source C：Stage D1 Sensitivity `run_20260917_01`；FORMAL_REFERENCE 只登记引用，不新增 canonical row。",
        "- Source D：Stage D2 MOVE Ablation `run_20260918_02`；COPY 只读引用 Formal，新增四个 MOVE row。", "",
        "所有 source manifest、checkpoint、progress、determinism、protocol、trace identity 与已冻结 provenance 均通过 gate。pilot、smoke、INVALID、SUPERSEDED_INCOMPLETE、O2/SR/P3 均未读取为数据源。", "",
        "## 2. Canonicalization 与去重", "",
        "Canonical key 使用 workload、routing、policy、action、N、theta、W、h、K、q 和 Evaluation 起止时间。family 不参与行为去重；D1/D2 的 Formal reference 通过 key 和原始指标交叉验证后写入 `referenced_by_families`。", "",
        f"最终母表共 **{row_counts['final_results_long']}** 行：Formal Main 14、D1 新 case 32、D2 MOVE 4。canonical ID 与行为 key 均无重复。", "",
        "## 3. Cohort 隔离与单位", "",
        "Main、theta、Persistence K、Recency q、N、MOVE 使用 `MAIN_25_45=[1500000,2700000)`；Oracle W sensitivity 使用 `W_SENS_10_25=[600000,1500000)`。只有 workload 和 cohort 完全相同才计算 delta。Saved/WeightedSaved 同时保留原值与 /1e6；wire 同时保留 bytes 与 decimal TB=/1e12。", "",
        "## 4. 输出表", "",
    ]
    for name, count in sorted(row_counts.items()):
        lines.append(f"- `{name}`：{count} rows。")
    lines += ["", "Excel 提供 README、主表、baseline、Oracle headroom、各 sensitivity、COPY/MOVE、Prompt buckets、Diagnostics、Strategy Master 与 All Long；CSV/JSON 是 canonical machine-readable source。", "",
              "## 5. 一致性检查", "", "| Check | Result |", "|---|---|"]
    for name, value in checks.items(): lines.append(f"| `{name}` | {'PASS' if value else 'FAIL'} |")
    lines += ["", "## 6. Provenance", "", "| Source | SHA-256 |", "|---|---|"]
    for name, value in source_sha.items(): lines.append(f"| `{name}` | `{value}` |")
    lines += ["", "最终输出 artifact SHA-256 记录在 `consolidation_manifest.json`。", "", "## 7. Missing / NA", "",
              "CSV 中 NA、JSON 中 null。MOVE-only release 字段对 COPY 为 NA；无 proactive transfer 的 unused ratio 为 NA；未由可靠 source 发出的 ordinary-LRU/churn 字段保持 NA，并用 scope/reason 字段说明。没有把未知值写成 0。", "",
              "## 8. 解释边界", "", "本阶段没有画图、没有补实验、没有参数选择，也没有写最终科研结论。WeightedSaved 不是延迟；unused transfer wire 不等于完全无 partial reuse；MOVE source release 与 target-side unused ratio 是不同指标。", "",
              "## 9. 最终状态", "", "**FINAL_RESULTS_CONSOLIDATION_PASS**", ""]
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")


def run(output: Path) -> None:
    gate = source_gate()
    d2_table_source = load_json(D2 / "tables/copy_vs_move.json")
    canonical = formal_rows(gate, d2_table_source)
    d1_new, d1_source = d1_rows(gate)
    canonical.extend(d1_new)
    canonical.extend(d2_rows(gate, d2_table_source))
    attach_reference_families(canonical, d1_source, d2_table_source)
    canonical.sort(key=lambda r: (r["workload"], r["evaluation_start_ms"], r["family"], r["line_id"] or 0, r["case_id"]))

    main_conversation, main_toolagent, main_combined = main_tables(canonical)
    baselines = baseline_table(canonical)
    oracle = oracle_headroom(canonical)
    history = history_summary(canonical)
    sensitivity = sensitivity_context(canonical, d1_source)
    sensitivity_tables = sensitivity_subtables(sensitivity)
    copy_move, copy_delta = copy_move_tables(d2_table_source)
    buckets, oracle_buckets = prompt_bucket_tables()
    diagnostics = diagnostics_table()
    strategy = strategy_master(main_combined, baselines, sensitivity, copy_move, canonical)

    checks = validate_all(canonical, main_combined, sensitivity, copy_move, copy_delta, buckets)

    if output.exists():
        raise RuntimeError(f"Output directory already exists: {output}")
    temp = output.with_name(output.name + ".tmp")
    if temp.exists(): shutil.rmtree(temp)
    temp.mkdir(parents=True)

    tables: dict[str, tuple[Sequence[Mapping[str, Any]], Sequence[str] | None]] = {
        "final_results_long": (canonical, FINAL_LONG_COLUMNS),
        "main_7line_conversation": (main_conversation, None),
        "main_7line_toolagent": (main_toolagent, None),
        "main_7line_combined": (main_combined, None),
        "baseline_comparison": (baselines, None), "oracle_headroom": (oracle, None),
        "history_strategy_summary": (history, None), "sensitivity_master": (sensitivity, None),
        **{name: (rows, None) for name, rows in sensitivity_tables.items()},
        "copy_vs_move": (copy_move, None), "copy_vs_move_delta": (copy_delta, None),
        "prompt_bucket_final": (buckets, None), "oracle_prompt_bucket_delta": (oracle_buckets, None),
        "diagnostics_summary": (diagnostics, None), "strategy_parameter_master": (strategy, None),
    }
    for name, (rows, columns) in tables.items(): write_pair(temp, name, rows, columns)
    data_dictionary(temp / "DATA_DICTIONARY.md")
    build_xlsx(temp / "strategy_parameter_master.xlsx", [
        ("Main_7Line", main_combined), ("Baselines", baselines), ("Oracle_Headroom", oracle),
        ("Persistence", sensitivity_tables["persistence_sensitivity_final"]),
        ("Recency", sensitivity_tables["recency_sensitivity_final"]),
        ("Oracle_W", sensitivity_tables["oracle_w_sensitivity_final"]),
        ("Theta", sensitivity_tables["theta_sensitivity_final"]),
        ("N", sensitivity_tables["n_sensitivity_final"]), ("COPY_vs_MOVE", copy_move),
        ("Prompt_Buckets", buckets), ("Diagnostics", diagnostics),
        ("Strategy_Master", strategy), ("All_Long", canonical),
    ])

    validation = {
        "status": "PASS", "final_state": "FINAL_RESULTS_CONSOLIDATION_PASS",
        "checks": checks, "source_gate": {"status": "PASS", "sha256": gate["source_sha256"]},
        "row_counts": {name: len(rows) for name, (rows, _) in tables.items()},
        "excluded_sources": ["pilot", "smoke", "INVALID", "SUPERSEDED_INCOMPLETE", "historical O2/SR/P3"],
        "simulator_invoked": False, "existing_results_modified": False,
    }
    write_json(temp / "consolidation_validation.json", validation)
    artifact_hashes = {str(p.relative_to(temp)): sha256(p) for p in sorted(temp.iterdir()) if p.is_file()}
    manifest = {
        "status": "PASS", "final_state": "FINAL_RESULTS_CONSOLIDATION_PASS",
        "protocol_version": PROTOCOL, "canonical_row_count": len(canonical),
        "source_sha256": gate["source_sha256"], "artifact_sha256": artifact_hashes,
        "validation_file": "consolidation_validation.json",
        "audit_report": "docs/analysis/task_main_final_results_consolidation.md",
    }
    write_json(temp / "consolidation_manifest.json", manifest)

    row_counts = validation["row_counts"]
    audit_report(AUDIT_REPORT, output, checks, row_counts, gate["source_sha256"], artifact_hashes)
    output.parent.mkdir(parents=True, exist_ok=True)
    temp.rename(output)
    print(json.dumps({"status": "FINAL_RESULTS_CONSOLIDATION_PASS", "output": str(output.relative_to(ROOT)),
                      "canonical_rows": len(canonical), "checks": len(checks),
                      "tables": len(tables)}, indent=2))


def main() -> None:
    parser = argparse.ArgumentParser(description="Consolidate verified TaskMain artifacts without simulation")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    output = args.output if args.output.is_absolute() else ROOT / args.output
    run(output)


if __name__ == "__main__":
    main()
